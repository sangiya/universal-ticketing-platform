#!/usr/bin/env python3
"""Generate the TicketMesh scope document in both .docx and .pdf formats.

Reads docs/spec/REQUIREMENTS.md (the single source of truth) plus a supplementary
executive summary, and renders a professional scope document. Run from the repo
root:  python tools/generate_docs.py
"""
import re
import subprocess
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
REQ = ROOT / "docs" / "spec" / "REQUIREMENTS.md"
OUT_DOCX = ROOT / "docs" / "TicketMesh_Scope_Document.docx"
OUT_PDF = ROOT / "docs" / "TicketMesh_Scope_Document.pdf"

EXEC_SUMMARY = (
    "TicketMesh is a universal, configurable, multi-provider ticketing and reservation "
    "SaaS that can be bought as software, self-hosted, or run as a full managed service "
    "with hosting. It uses a marketplace (Uber/PickMe-style) model where shops and agents "
    "connect via app, upload and sell their services, and customers also self-serve. The "
    "product is global and white-label themeable (country, currency, language, logo, "
    "colors — all by configuration, no code), delivered across a consumer Web channel "
    "and a mobile-responsive PWA app channel, plus an admin portal. It ships with 24/7 "
    "support, fraud detection, auto issue detection and auto-fix, full SDLC documentation, "
    "test and production API docs, and deploys to AWS, Kubernetes/EKS, on-premises and "
    "Docker with horizontal autoscaling."
)


def add_code(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"[CODE:\1]", text)
    return text


def markdown_to_runs(paragraph, text):
    """Minimal renderer: bold **...**, inline code `...` > [CODE:...]."""
    parts = re.split(r"(\*\*.*?\*\*|\[CODE:.*?\])", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("[CODE:"):
            code = part[len("[CODE:"):-1]
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("TicketMesh", level=0)
    sub = doc.add_paragraph("Universal Configurable Ticketing & Reservation Platform — "
                            "Single Source of Truth Scope Document")
    sub.runs[0].italic = True
    for p in doc.paragraphs[-1:]:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    doc.add_paragraph("Executive Summary").style = doc.styles["Heading 1"]
    doc.add_paragraph(EXEC_SUMMARY)

    lines = REQ.read_text(encoding="utf-8").splitlines()
    in_code_fence = False
    for raw in lines:
        stripped = raw.rstrip()
        if stripped.startswith("```"):
            in_code_fence = not in_code_fence
            continue
        if stripped.startswith("# ") and not in_code_fence:
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## ") and not in_code_fence:
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### ") and not in_code_fence:
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("- ") and not in_code_fence:
            p = doc.add_paragraph(style="List Bullet")
            markdown_to_runs(p, add_code(stripped[2:].strip()))
        elif stripped.strip() and not in_code_fence:
            p = doc.add_paragraph()
            markdown_to_runs(p, add_code(stripped.strip()))
        elif not in_code_fence:
            pass

    doc.add_page_break()
    doc.add_heading("Author", level=1)
    doc.add_paragraph("Prepared by sangiya — Associate Software Architect.")
    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX}")


def build_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer,
        PageBreak, ListFlowable, ListItem,
    )
    from reportlab.lib import colors

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16, spaceAfter=8,
                        textColor=colors.HexColor("#0B3B60"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceBefore=8,
                        spaceAfter=4, textColor=colors.HexColor("#0B3B60"))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=12, bulletIndent=6, spaceAfter=2)

    def esc(t):
        t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        t = t.replace("`", "")
        return t

    story = []
    story.append(Paragraph("TicketMesh", h1))
    story.append(Paragraph(
        "Universal Configurable Ticketing &amp; Reservation Platform — Scope Document",
        ParagraphStyle("sub", parent=body, fontSize=11, alignment=TA_CENTER, textColor=colors.grey)))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Executive Summary", h2))
    story.append(Paragraph(esc(EXEC_SUMMARY), body))
    story.append(PageBreak())

    lines = REQ.read_text(encoding="utf-8").splitlines()
    in_fence = False
    for raw in lines:
        s = raw.rstrip()
        if s.startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        if s.startswith("# "):
            story.append(Paragraph(esc(s[2:].strip()), h1))
        elif s.startswith("## "):
            story.append(Paragraph(esc(s[3:].strip()), h2))
        elif s.startswith("### "):
            story.append(Paragraph(esc(s[4:].strip()), h2))
        elif s.startswith("- "):
            story.append(ListFlowable(
                [ListItem(Paragraph(esc(s[2:].strip()), body), leftIndent=10)],
                bulletType="bullet", start="•"))
        elif s.strip():
            story.append(Paragraph(esc(s.strip()), body))

    story.append(PageBreak())
    story.append(Paragraph("Author", h1))
    story.append(Paragraph("Prepared by sangiya — Associate Software Architect.", body))

    SimpleDocTemplate(str(OUT_PDF), pagesize=A4,
                      leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm).build(story)
    print(f"Wrote {OUT_PDF}")


def main():
    build_docx()
    build_pdf()
    print("Scope document generation complete.")


if __name__ == "__main__":
    sys.exit(main())
